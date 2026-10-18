"""Document text extraction parsers for PDF, DOCX, TXT, and chat transcripts.

Specification Section 8.1.
Supports:
- PDF: pypdf (raises ValueError for scanned PDFs with no extractable text)
- DOCX: python-docx (extracts paragraph and table text)
- TXT: UTF-8 plain text reader
- Chat transcripts: Formatted turn extraction
"""

from typing import List
from pypdf import PdfReader
import docx


def parse_pdf(file_path: str) -> str:
    """Extracts text from a PDF file.

    Args:
        file_path: Path to the PDF file.

    Returns:
        Extracted raw text.

    Raises:
        ValueError: If the PDF is scanned or has no extractable text.
    """
    reader = PdfReader(file_path)
    extracted_text = []

    for page_num, page in enumerate(reader.pages):
        text = page.extract_text()
        if text:
            extracted_text.append(text)

    full_text = "\n".join(extracted_text).strip()
    if not full_text:
        raise ValueError("Scanned PDF not supported: No extractable text found in document.")

    return full_text


def parse_docx(file_path: str) -> str:
    """Extracts text from a DOCX file including paragraphs and tables.

    Args:
        file_path: Path to the DOCX file.

    Returns:
        Extracted raw text.
    """
    doc = docx.Document(file_path)
    text_chunks = []

    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            text_chunks.append(paragraph.text.strip())

    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                text_chunks.append(row_text)

    return "\n".join(text_chunks)


def parse_txt(file_path: str) -> str:
    """Reads text from a plain text file with UTF-8 or latin-1 fallback.

    Args:
        file_path: Path to the TXT file.

    Returns:
        Content string.
    """
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            return f.read()
    except UnicodeDecodeError:
        with open(file_path, "r", encoding="latin-1") as f:
            return f.read()


def format_chat_transcript(messages: List[dict], chat_date: str) -> str:
    """Formats student chat turns into clean readable transcript text for RAG indexing.

    Args:
        messages: List of message dicts containing 'role' and 'content'.
        chat_date: Date string string for header.

    Returns:
        Formatted transcript string.
    """
    lines = [f"--- Chat Transcript ({chat_date}) ---"]
    for msg in messages:
        role = "Student" if msg.get("role") == "user" else "Assistant"
        lines.append(f"{role}: {msg.get('content', '').strip()}")
    return "\n".join(lines)
